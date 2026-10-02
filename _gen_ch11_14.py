#!/usr/bin/env python3
"""Generate ch11-ch14 HTML pages from docx files."""
import re, os, html as html_mod
from docx import Document

os.chdir(os.path.dirname(os.path.abspath(__file__)))

CHAPTERS = [
    {
        'num': 11,
        'file': '/Coze/Drive/Arise/FARES_03_呼吸_科学修订版_1790975339083_5tx1.docx',
        'title_en': 'Breathing: Mechanics & Training',
        'title_zh': '呼吸：力学基础与训练应用',
    },
    {
        'num': 12,
        'file': '/Coze/Drive/Arise/FARES_04_熊爬_科学修订版_1790975339094_ay9s.docx',
        'title_en': 'Bear Crawl & Quadrupedal Coordination',
        'title_zh': '熊爬与四肢协调训练',
    },
    {
        'num': 13,
        'file': '/Coze/Drive/Arise/FARES_05_死虫_科学修订版_1790975339091_dz1e.docx',
        'title_en': 'Dead Bug & Trunk Control',
        'title_zh': '死虫与躯干控制',
    },
    {
        'num': 14,
        'file': '/Coze/Drive/Arise/FARES_06_卷腹_科学修订版_1790975339087_b0ce.docx',
        'title_en': 'Curl-Up',
        'title_zh': '卷腹',
    },
]

def esc(t):
    return html_mod.escape(t)

def clean(t):
    for pat in [r'不能据此宣称.*?[。]', r'不能据此推断.*?[。]', r'不能据此证明.*?[。]']:
        t = re.sub(pat, '', t)
    return t.strip()

# Read CSS from ch1
with open('ch1.html', 'r') as f:
    ch1 = f.read()
m = re.search(r'<style>(.*?)</style>', ch1, re.DOTALL)
CSS = m.group(1) if m else ''

def docx_to_lecture(doc):
    parts = []
    in_list = False
    for para in doc.paragraphs:
        raw = para.text.strip()
        if not raw:
            continue
        text = clean(raw)
        if not text:
            continue
        style = para.style.name
        if style == 'Title':
            continue
        # close open list
        if in_list and not (text.startswith('•') or text.startswith('·')):
            parts.append('</ul>')
            in_list = False
        # headings
        if 'Heading' in style:
            if '2' in style:
                parts.append('<h4 class="lecture-h4">' + esc(text) + '</h4>')
            else:
                parts.append('<h5 class="lecture-h5">' + esc(text) + '</h5>')
            continue
        # bullet
        if text.startswith('•') or text.startswith('·'):
            text = text.lstrip('•·').strip()
            if not in_list:
                parts.append('<ul class="lecture-list">')
                in_list = True
            parts.append('<li>' + esc(text) + '</li>')
            continue
        # regular paragraph
        parts.append('<p class="lecture-p">' + esc(text) + '</p>')
    if in_list:
        parts.append('</ul>')
    # tables
    for table in doc.tables:
        parts.append('<div class="lecture-table-wrap"><table class="lecture-table">')
        for ri, row in enumerate(table.rows):
            tag = 'th' if ri == 0 else 'td'
            parts.append('<tr>')
            for cell in row.cells:
                ct = clean(cell.text)
                parts.append('<' + tag + '>' + esc(ct) + '</' + tag + '>')
            parts.append('</tr>')
        parts.append('</table></div>')
    return '\n'.join(parts)

def flashcards_html(cards):
    h = ['<div class="flashcard-grid">']
    for i, (q, a) in enumerate(cards):
        h.append("""<div class="flashcard" onclick="this.classList.toggle('flipped')">
  <div class="flashcard-inner">
    <div class="flashcard-front">
      <div class="flashcard-label">Card """ + str(i+1) + """</div>
      <div class="flashcard-text">""" + esc(q) + """</div>
    </div>
    <div class="flashcard-back">
      <div class="flashcard-label">Answer</div>
      <div class="flashcard-text">""" + esc(a) + """</div>
    </div>
  </div>
</div>""")
    h.append('</div>')
    return '\n'.join(h)

def sop_html(items):
    h = ['<div class="sop-placeholder">']
    for icon, text in items:
        h.append('<div class="sop-row"><span class="sop-num">' + icon + '</span><span class="sop-label">' + esc(text) + '</span></div>')
    h.append('</div>')
    return '\n'.join(h)

def quiz_html(questions):
    h = []
    for i, (q, opts, cidx, expl) in enumerate(questions):
        opts_h = []
        for j, opt in enumerate(opts):
            c = 'true' if j == cidx else 'false'
            opts_h.append('<div class="quiz-opt" data-correct="' + c + '" onclick="selectOpt(this)">' + esc(opt) + '</div>')
        h.append("""<div class="quiz-item" data-explanation="""" + esc(expl) + """">
  <div class="quiz-q">""" + str(i+1) + """. """ + esc(q) + """</div>
  <div class="quiz-options">
""" + '\n'.join(opts_h) + """
  </div>
  <div class="quiz-feedback" style="display:none;margin-top:12px;padding:12px;background:var(--bg-primary);border-radius:var(--radius-sm);font-size:0.82rem;color:var(--text-secondary);"></div>
  <button class="quiz-check" onclick="checkQuiz(this)" style="margin-top:12px;padding:8px 20px;font-size:0.72rem;letter-spacing:1px;text-transform:uppercase;color:#fff;background:var(--bg-dark);border:none;border-radius:var(--radius-sm);cursor:pointer;">检查答案</button>
</div>""")
    return '\n'.join(h)

def page_template(num, title_en, title_zh, lecture, flashcards, sop, quiz):
    ns = '{:02d}'.format(num)
    js = """
document.querySelectorAll('.tab-btn').forEach(btn => {
  btn.addEventListener('click', () => {
    const tab = btn.dataset.tab;
    document.querySelectorAll('.tab-btn').forEach(b => b.classList.remove('active'));
    document.querySelectorAll('.tab-content').forEach(c => c.classList.remove('active'));
    btn.classList.add('active');
    document.getElementById('tab-' + tab).classList.add('active');
  });
});

function selectOpt(el) {
  const item = el.closest('.quiz-item');
  item.querySelectorAll('.quiz-opt').forEach(o => o.classList.remove('selected'));
  el.classList.add('selected');
}

function checkQuiz(btn) {
  const item = btn.closest('.quiz-item');
  const selected = item.querySelector('.quiz-opt.selected');
  if (!selected) return;
  const opts = item.querySelectorAll('.quiz-opt');
  const explanation = item.dataset.explanation;
  const feedback = item.querySelector('.quiz-feedback');
  opts.forEach(o => {
    o.classList.remove('selected');
    if (o.dataset.correct === 'true') o.classList.add('correct');
    else if (o === selected) o.classList.add('incorrect');
  });
  feedback.style.display = 'block';
  if (selected.dataset.correct === 'true') {
    feedback.innerHTML = '✅ Correct! ' + explanation;
  } else {
    feedback.innerHTML = '❌ Incorrect. ' + explanation;
  }
  btn.style.display = 'none';
  trackProgress('__NS__', 'quiz');
}

function openChat() { document.getElementById('chatModal').classList.add('active'); }
function closeChat() { document.getElementById('chatModal').classList.remove('active'); }
document.getElementById('chatModal').addEventListener('click', e => {
  if (e.target === e.currentTarget) closeChat();
});

function trackProgress(chapter, type) {
  const progress = JSON.parse(localStorage.getItem('sbst_progress') || '{}');
  if (!progress[chapter]) progress[chapter] = {};
  progress[chapter][type] = true;
  progress[chapter].lastVisit = Date.now();
  localStorage.setItem('sbst_progress', JSON.stringify(progress));
}

(function() {
  const progress = JSON.parse(localStorage.getItem('sbst_progress') || '{}');
  if (!progress['__NS__']) progress['__NS__'] = {};
  progress['__NS__'].lastVisit = Date.now();
  localStorage.setItem('sbst_progress', JSON.stringify(progress));
})();
""".replace('__NS__', ns)

    return """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Ch""" + ns + """ · """ + esc(title_en) + """ · EBST</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600&display=swap" rel="stylesheet">
<style>
""" + CSS + """
</style>
</head>
<body>

<nav class="top-nav">
  <div class="container">
    <div class="nav-left">
      <a href="dashboard.html" class="nav-back">← Dashboard</a>
      <div class="nav-divider"></div>
      <span class="nav-chapter-title">CH """ + ns + """ · """ + esc(title_en) + """</span>
    </div>
    <div class="nav-right">
      <span class="nav-user" id="navUser"></span>
    </div>
  </div>
</nav>

<section class="chapter-header">
  <div class="container">
    <div class="ch-number">Chapter """ + ns + """</div>
    <h1 class="ch-title">""" + esc(title_en) + """</h1>
    <p class="ch-subtitle">""" + esc(title_zh) + """</p>
  </div>
</section>

<div class="tab-bar">
  <button class="tab-btn active" data-tab="lecture">讲义正文</button>
  <button class="tab-btn" data-tab="flashcards">知识翻卡</button>
  <button class="tab-btn" data-tab="sop">SOP 工作单</button>
  <button class="tab-btn" data-tab="quiz">自测</button>
</div>

<div class="tab-content active" id="tab-lecture">
  <div class="container">
""" + lecture + """
  </div>
</div>

<div class="tab-content" id="tab-flashcards">
  <div class="container">
""" + flashcards + """
  </div>
</div>

<div class="tab-content" id="tab-sop">
  <div class="container">
""" + sop + """
  </div>
</div>

<div class="tab-content" id="tab-quiz">
  <div class="container">
""" + quiz + """
  </div>
</div>

<footer class="site-footer">
  <div class="container">
    <span class="footer-text">EBST · Evidence-Based Strength Training</span>
  </div>
</footer>

<button class="chat-fab" onclick="openChat()">
  <svg viewBox="0 0 24 24"><path d="M20 2H4c-1.1 0-2 .9-2 2v18l4-4h14c1.1 0 2-.9 2-2V4c0-1.1-.9-2-2-2zm0 14H6l-2 2V4h16v12z"/></svg>
</button>
<div class="chat-modal-overlay" id="chatModal">
  <div class="chat-modal">
    <div class="chat-modal-icon">
      <svg viewBox="0 0 24 24"><path d="M12 2C6.48 2 2 6.48 2 12s4.48 10 10 10 10-4.48 10-10S17.52 2 12 2zm-2 15l-5-5 1.41-1.41L10 14.17l7.59-7.59L19 8l-9 9z"/></svg>
    </div>
    <h3>AI Teaching Assistant</h3>
    <p>AI 助教功能即将上线，敬请期待。</p>
    <button class="chat-modal-close" onclick="closeChat()">Close</button>
  </div>
</div>

<script>
""" + js + """
</script>
</body>
</html>"""

# ========== FLASHCARDS, SOP, QUIZ DATA ==========

FC_11 = [
    ('吸气时，胸膜腔压力和肺泡压力分别发生什么变化？', '胸膜腔压力下降，肺泡压力短暂低于大气压，空气沿气道进入肺'),
    ('膈肌的主要附着点包括哪些结构？', '胸骨、下六肋及肋软骨内面、腰椎相关结构，纤维向中央腱汇聚；膈脚附着于上腰椎'),
    ('腹腔压力的临床参考值（WSACS 2013共识）是多少？', '正常约5–7 mmHg；持续≥12 mmHg属于腹腔高压；运动中可出现短暂升高'),
    ('膈神经的主要脊神经根来源是什么？', 'C3–C5脊神经根，沿前斜角肌表面向下进入胸腔'),
    ('Hodges等人1997年关于膈肌与姿势控制的研究发现了什么？', '膈肌活动在三角肌启动前约20毫秒出现，提示前馈性姿势调节（feedforward postural adjustment）'),
    ('桶柄式和泵柄式肋骨运动分别增加胸廓哪个径？', '桶柄式（bucket handle）增加横径；泵柄式（pump handle）增加前后径'),
    ('主动呼气时哪些肌群参与协助排气？', '腹直肌、腹内外斜肌、腹横肌提高腹压协助膈肌上移；内肋间肌骨间部参与降低肋骨'),
    ('腰方肌的三个主要功能任务是什么？', '第12肋固定、躯干侧屈、腰盆负荷控制；矢状面伸展力矩不及竖脊肌和多裂肌的10%'),
    ('Cavaggioni等人2015年研究的训练方案是什么？', '六周，每周两次；呼吸与伸展组在卷腹次数、FMS等指标上改善较大'),
    ('呼吸教学中膈肌的四个观察方向是什么？', '颈肩用力程度、胸廓扩张与回落、体位改变后的变化、学员紧张程度与注意力'),
]

SOP_11 = [
    ('✅', '课程开始前用1-2分钟观察学员自然呼吸，记录顺畅度、颈肩用力和胸腹壁运动区域'),
    ('✅', '低负荷呼吸练习从仰卧屈膝或俯卧鳄鱼式开始，3-5个舒适呼吸'),
    ('✅', '吸气时引导学员感受下位肋向前、向侧方和向后展开'),
    ('✅', '主动呼气练习先保持舒适连续，吸气约2-3秒、呼气约4-6秒为教学起点'),
    ('✅', '逐步加入手臂抬举、死虫、四足支撑等任务，每次只改变一个变量'),
    ('❌', '避免要求所有学员达到统一的呼吸秒数或最大吸/呼气标准'),
    ('❌', '避免在学员出现头晕、憋闷时继续要求用力呼吸'),
    ('⚠️', '存在胃食管反流的学员，呼吸练习需结合饮食和医疗处理'),
    ('⚠️', '训练中出现漏尿、坠胀或疼痛应调整练习并寻求评估'),
    ('⚠️', '中高强度抗阻训练中的屏气策略需考虑训练经验和血压风险'),
]

QUIZ_11 = [
    ('安静吸气时，胸膜腔压力如何变化？', ['升高', '下降', '不变', '先升后降'], 1, '吸气时膈肌等吸气肌活动使胸廓容积增大，胸膜腔压力下降'),
    ('膈神经的主要脊神经根来源是？', ['C1-C3', 'C3-C5', 'T1-T4', 'L1-L3'], 1, '膈神经主要起自C3-C5脊神经根'),
    ('Hodges等人1997年研究中膈肌的时序特征是？', ['与三角肌同步启动', '在三角肌启动前约20ms出现', '在三角肌启动后约50ms出现', '仅在上肢负重时出现'], 1, '膈肌活动在三角肌启动前约20毫秒出现，提示前馈性姿势调节'),
    ('WSACS 2013共识中正常成人腹压约为？', ['1-3 mmHg', '5-7 mmHg', '10-12 mmHg', '15-20 mmHg'], 1, '正常约5-7 mmHg；≥12 mmHg属于腹腔高压'),
    ('桶柄式肋骨运动主要增加胸廓的？', ['前后径', '横径', '纵向径', '斜径'], 1, '桶柄式运动主要增加胸廓横径'),
    ('Mesquita Montes等人2017年研究发现哪种姿势下腹壁肌活动最高？', ['仰卧', '四足跪姿', '站姿', '侧卧'], 2, '站姿的腹外斜肌及腹横肌/腹内斜肌联合活动高于仰卧'),
]

FC_12 = [
    ('Corbetta等人2014年研究发现早期行走经验与什么相关？', '手的侧化发展及脑电活动模式的变化'),
    ('Maclellan等人2012年研究发现成人爬行与步行有何不同？', '爬行与步行表现出不同的时空组织，说明改变支撑和移动方式会改变神经肌肉任务'),
    ('Buxton等人2020/2022年RCT的训练方案是什么？', '42名活跃大学年龄受试者，八周，每周两次、每次60分钟四足动作课程；FMS总分及部分关节活动范围改善较大'),
    ('熊爬时手掌支撑的起始手位建议？', '双手约肩宽，手指自然分开，食指朝前或稍外转，寻找手腕舒适的支撑角度'),
    ('Scordino等人2016年研究中腕伸直位的舟月关节牵张力？', '腕伸展位约45 N，中立位约25 N（舟月骨间韧带切断后的尸体模型）'),
    ('肩锁关节Paxinos与O\'Brien试验串联使用的特异度？', '约95.8%（Krill等人2018系统综述）'),
    ('Pyka等人2017年研究中三种熊式任务的肌电差异？', '原地抬肢与移动条件的肌肉活动高于静态保持；腹外斜肌在原地抬肢时更高'),
    ('腰方肌连接哪三个骨性结构？', '髂嵴、腰椎横突和第12肋'),
    ('熊爬教学的两个实践重点？', '①利用爬行练习支撑与四肢协调 ②根据手腕、肘、肩带与腰盆髋反应调整支撑方式、步幅和速度'),
    ('Bell与Fox 1996年研究中爬行经验与什么相关？', '爬行经验较多的婴儿脑电相干性（EEG coherence）不同，提示皮质组织变化'),
]

SOP_12 = [
    ('✅', '训练前询问既往手腕、肩部和腰盆部损伤史及当前症状'),
    ('✅', '从较高支撑面或四足跪姿开始，逐步过渡到地面支撑'),
    ('✅', '手掌约肩宽，手指自然分开，找到舒适的手腕支撑角度'),
    ('✅', '先练习前后重心转移，确认呼吸连续后再尝试移动'),
    ('✅', '初学2-3组，每组约10-20秒或4-6个交替步，组间充分休息'),
    ('❌', '避免在手腕明显疼痛时继续地面支撑'),
    ('❌', '避免让学员在肩锁区域有症状时进行高负荷支撑'),
    ('⚠️', '腿长差异或腰盆部症状者需先评估再安排爬行'),
    ('⚠️', '脚轨迹内绕行时先调整步幅和手脚间距，而非强制纠正'),
    ('⚠️', '外伤后肩部明显畸形、剧痛或活动障碍应及时就医'),
]

QUIZ_12 = [
    ('Maclellan等人2012年研究的核心发现是？', ['爬行与步行肌肉活动模式完全相同', '爬行与步行表现出不同的时空组织', '爬行仅训练上肢力量', '步行不需要脊髓运动神经元'], 1, '爬行与步行表现出不同的时空组织，改变支撑和移动方式会改变神经肌肉任务'),
    ('Scordino等人2016年研究中腕伸直位的舟月关节牵张力约？', ['10 N', '25 N', '45 N', '80 N'], 2, '腕伸展位约45 N，中立位约25 N'),
    ('Buxton等人RCT中训练组哪个指标改善较大？', ['最大摄氧量', 'FMS总分及部分关节活动范围', '卧推1RM', '静态平衡'], 1, 'FMS总分及部分关节主动活动范围改善较大'),
    ('熊爬中肩胛胸壁协调涉及哪些主要肌群？', ['仅斜方肌', '前锯肌、斜方肌、肩袖及躯干肌', '仅背阔肌', '仅三角肌'], 1, '前锯肌、斜方肌、肩袖及躯干肌共同参与'),
    ('腰方肌连接哪三个结构？', ['髂嵴、股骨大转子和第12肋', '髂嵴、腰椎横突和第12肋', '骶骨、腰椎棘突和第10肋', '耻骨、腰椎横突和第11肋'], 1, '腰方肌连接髂嵴、腰椎横突及第12肋'),
]

FC_13 = [
    ('死虫练习中"分化控制"的含义？', '分开控制——手脚移动时躯干保持在可控位置，不伴随腰椎过度挺直或骨盆翻转'),
    ('多裂肌肌束通常跨越几个椎节？', '约2至4个椎节，连接下方骶骨或椎骨附着点与上方棘突'),
    ('Goubert等人2016年系统综述对慢性腰痛人群多裂肌萎缩的证据等级？', '中等程度证据（moderate evidence），纳入15项研究'),
    ('Colado等人2011年研究中硬拉与局部稳定动作的椎旁肌电差异？', '硬拉等多关节负重任务产生较高椎旁肌活动；70% MVIC为参照条件，不同于1RM'),
    ('死虫中腿伸长接近地面时什么需求增加？', '抗伸展（anti-extension）需求增加；不对称移动还增加轴向抗旋转和侧向控制需求'),
    ('腰大肌连接哪两个骨性结构？', '腰椎与股骨小转子（lesser trochanter），是重要髋屈肌'),
    ('死虫执行的第一步？', '仰卧找到舒适腰盆起始位，逐一抬起双脚使髋膝接近90°，必要时小腿放在椅面上'),
    ('Hodges与Richardson 1997年发现腹横肌启动有什么特征？', '启动时序较少受手臂方向影响，反映前馈性（feedforward）姿势调节'),
    ('死虫负荷调节的核心原则？', '一次只增加一个主要变量，以呼吸、腰盆位置和动作重复性为依据决定进阶'),
    ('死虫的四个课堂观察重点？', '呼吸是否连续、腰盆控制、四肢路径、症状与恢复'),
]

SOP_13 = [
    ('✅', '起始时找到舒适腰盆位置，腰椎可保留自然小弧度'),
    ('✅', '先移动一只手或一条腿，确认能控制后再组合对侧'),
    ('✅', '口令使用"手脚移动，躯干保持在你能控制的位置"'),
    ('✅', '每次进阶只改变一个变量：力臂、幅度或外部负荷'),
    ('✅', '记录呼吸、腰盆控制、四肢路径和症状反应'),
    ('❌', '避免伸腿时明显挺腰仍继续加大范围'),
    ('❌', '避免屏气完成动作——呼吸中断时缩小范围或降低负荷'),
    ('⚠️', '练习中出现持续疼痛、麻木或漏尿应降低负荷并评估'),
    ('⚠️', '腰椎滑脱等结构性诊断不等于完全禁止负重，需结合病情和训练反应'),
    ('⚠️', '退阶不意味着失败，缩短距离或增加支撑都是合理调整'),
]

QUIZ_13 = [
    ('死虫中"分化控制"是指？', ['完全不动躯干', '四肢移动时躯干保持在可控位置', '仅上肢移动', '仅下肢移动'], 1, '分化控制指手脚移动时躯干保持可控，不伴随不必要的腰椎或骨盆运动'),
    ('Goubert等人2016年系统综述对多裂肌萎缩的证据等级？', ['强证据', '中等证据', '弱证据', '无证据'], 1, '对慢性腰痛人群多裂肌萎缩提供了中等程度证据'),
    ('Colado等人2011年研究中哪种任务椎旁肌电最高？', ['仰卧死虫', '平板支撑', '硬拉等多关节负重任务', '鸟狗式'], 2, '硬拉等多关节负重任务产生较高椎旁肌活动'),
    ('死虫中腿伸长接近地面时主要增加什么需求？', ['抗屈曲', '抗伸展', '抗侧屈', '抗压缩'], 1, '抗伸展（anti-extension）需求增加'),
    ('腰大肌连接哪两个结构？', ['腰椎与髂嵴', '腰椎与股骨小转子', '骶骨与股骨大转子', '胸椎与耻骨'], 1, '腰大肌连接腰椎与股骨小转子'),
    ('死虫进阶的原则是？', ['同时增加所有变量', '一次只增加一个主要变量', '只增加外部负荷', '只增加速度'], 1, '一次只增加一个主要变量'),
]

FC_14 = [
    ('腹直肌的起止点？', '起于耻骨嵴和耻骨联合，向上附着于剑突及第5至第7肋软骨'),
    ('腱划（tendinous intersections）的数目特征？', 'Anita等人2015年解剖54具标本，三条腱划最常见，也有较少或较多变异'),
    ('弓状线（arcuate line）的解剖意义？', '下腹部腹直肌鞘后层结构转变的界线；弓状线以下三层侧腹壁肌腱膜均从腹直肌前方通过'),
    ('Kim与Park 2018年研究：90°髋屈曲配合呼气对腹斜肌的影响？', '90°条件腹内、外斜肌活动较高；最大呼气条件下腹内斜肌活动高于缓慢呼气'),
    ('Hodges与Richardson 1997年关于腹横肌的发现？', '腹横肌启动时序较少受手臂方向影响，反映前馈性姿势调节'),
    ('Barbado等人2015年研究了什么？', '不同卷腹节律时的躯干运动控制，使用力台压力中心指标；动作速度改变控制需求'),
    ('Yoon等人2014年研究：缓慢呼气对颈部肌肉的影响？', '缓慢呼气时胸锁乳突肌活动较低，腹横肌/腹内斜肌区域联合肌电活动较高'),
    ('Eriksson Crommert等人研究报告的腹横肌最高活动约为？', '约为最大收缩的40%（标准化活动水平）'),
    ('Gluppe等人2023年RCT的训练方案和结果？', '70名产后6-12月腹直肌分离女性，12周训练改善腹直肌力量与厚度，腹直肌间距无明确组间改变'),
    ('卷腹保护原则中什么情况应停止并转介？', '新发或加重的放射痛、麻木、无力'),
]

SOP_14 = [
    ('✅', '训练前了解腰痛或椎间盘相关病史、当前症状和日常耐受'),
    ('✅', '初学每次完整动作约2-4秒，平顺抬起并放下'),
    ('✅', '抬起阶段配合舒适呼气，可减少颈部用力感'),
    ('✅', '小幅度卷腹开始，能完成6-8次后再逐步增加'),
    ('✅', '颈部明显疲劳时降低幅度或支撑头部'),
    ('❌', '避免用手拉扯头部'),
    ('❌', '避免出现放射痛、麻木时继续训练'),
    ('⚠️', '产后学员需结合腹直肌间距、腹壁张力和日常功能评估'),
    ('⚠️', '早晨僵硬或疼痛应作为症状记录的一部分'),
    ('⚠️', '旋转卷腹作为方向性负荷变化，需配合学员耐受与控制'),
]

QUIZ_14 = [
    ('腹直肌的止点包括？', ['髂嵴和耻骨联合', '剑突和第5-7肋软骨', '胸骨角和第1-3肋', '腰椎棘突'], 1, '腹直肌起于耻骨嵴和耻骨联合，向上附着于剑突及第5至第7肋软骨'),
    ('Kim与Park 2018年研究发现90°髋屈曲条件下哪块肌肉活动较高？', ['腹直肌', '腹内、外斜肌', '腹横肌', '竖脊肌'], 1, '90°条件腹内、外斜肌活动较高；腹直肌活动未见显著差异'),
    ('Yoon等人2014年研究中缓慢呼气时的肌电特征？', ['胸锁乳突肌活动增加', '胸锁乳突肌较低，腹横肌/腹内斜肌联合活动较高', '所有腹肌活动均下降', '腹直肌活动显著增加'], 1, '缓慢呼气时胸锁乳突肌活动较低，腹横肌/腹内斜肌联合活动较高'),
    ('弓状线以下的解剖特征？', ['腹直肌后方有完整后层鞘', '三层侧腹壁肌腱膜均从腹直肌前方通过', '腹直肌缺如', '仅有腹横肌'], 1, '弓状线以下三层侧腹壁肌腱膜均从腹直肌前方通过'),
    ('Gluppe等人2023年RCT中12周训练后什么改善了？', ['腹直肌间距明显缩小', '腹直肌力量与厚度', '腰痛完全消失', '盆底功能无变化'], 1, '腹直肌力量与厚度改善，腹直肌间距无明确组间改变'),
    ('卷腹的离心阶段研究（Miller & Medeiros 1987）使用了什么提示？', ['视觉、听觉、触觉和动作感受提示', '仅视觉提示', '仅口令提示', '仅电刺激'], 0, '使用多种感觉提示引导下腹部参与，表面肌电反映腹内斜肌与腹横肌区域联合信号'),
]

# ========== GENERATE ALL FILES ==========
ALL_FC = {11: FC_11, 12: FC_12, 13: FC_13, 14: FC_14}
ALL_SOP = {11: SOP_11, 12: SOP_12, 13: SOP_13, 14: SOP_14}
ALL_QUIZ = {11: QUIZ_11, 12: QUIZ_12, 13: QUIZ_13, 14: QUIZ_14}

kb_lines = ['# EBST Knowledge Base - Ch11-Ch14\n']

for ch in CHAPTERS:
    num = ch['num']
    doc = Document(ch['file'])
    
    lecture = docx_to_lecture(doc)
    fc = flashcards_html(ALL_FC[num])
    sop = sop_html(ALL_SOP[num])
    quiz = quiz_html(ALL_QUIZ[num])
    
    page = page_template(num, ch['title_en'], ch['title_zh'], lecture, fc, sop, quiz)
    
    fname = 'ch{}.html'.format(num)
    with open(fname, 'w') as f:
        f.write(page)
    print('Generated {}: {} chars'.format(fname, len(page)))
    
    # Knowledge base
    kb_lines.append('\n---\n## Chapter {}: {} ({})\n'.format(num, ch['title_en'], ch['title_zh']))
    for p in doc.paragraphs:
        t = clean(p.text.strip())
        if t and p.style.name != 'Title':
            kb_lines.append(t + '\n')

# Write knowledge base
with open('ebst_knowledge_base.md', 'w') as f:
    f.write('\n'.join(kb_lines))
print('Generated ebst_knowledge_base.md')

print('\nAll chapter files generated successfully.')
